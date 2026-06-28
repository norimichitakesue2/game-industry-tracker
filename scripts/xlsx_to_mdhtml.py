#!/usr/bin/env python3
"""
xlsx → md/html 変換スクリプト
リポジトリルートで `python scripts/xlsx_to_mdhtml.py` 実行
"""
import os
import sys
import html
import re
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = REPO_ROOT / "master" / "ゲーム業界・時事情報収集.xlsx"
MD_DIR = REPO_ROOT / "md"
DOCS_DIR = REPO_ROOT / "docs"

SHEET_SLUGS = {
    "README": "00-readme",
    "日次ニュース": "01-daily-news",
    "新作リリースログ": "02-releases",
    "IP・タイトルウォッチ": "03-ip-watch",
    "企業分析（国内）": "04-companies-jp",
    "企業分析（海外大手）": "05-companies-overseas",
    "四半期決算サマリ": "06-earnings",
    "業界トレンド・技術": "07-trends",
    "インディー注目": "08-indie",
    "今後起きそうなこと": "09-future",
    "掘り下げ・分析": "10-deep-dive",
    "業界マトリクス": "11-matrix",
    "横断的問い": "12-cross-questions",
    "週間ランキング": "13-rankings",
    "イベントカレンダー": "14-events",
    "ゲームデザイン・トレンド": "15-game-design",
}

# 「行ごとセクション形式」で表示するシート
SECTION_SHEETS = {"日次ニュース", "今後起きそうなこと", "掘り下げ・分析", "横断的問い"}

CSS = """
<style>
:root{
  --bg:#0d1117; --panel:#161b22; --border:#30363d;
  --text:#c9d1d9; --muted:#8b949e; --accent:#58a6ff; --accent-2:#79c0ff;
  --header-bg:#1f3864; --header-text:#fff;
}
@media (prefers-color-scheme: light){
  :root{ --bg:#f6f8fa; --panel:#fff; --border:#d0d7de;
         --text:#1f2328; --muted:#656d76; --accent:#0969da; --accent-2:#0550ae;
         --header-bg:#1f3864; --header-text:#fff; }
}
*{box-sizing:border-box}
body{
  background:var(--bg); color:var(--text);
  font-family:-apple-system,BlinkMacSystemFont,"Helvetica Neue","Hiragino Sans","Yu Gothic","Noto Sans JP",sans-serif;
  margin:0; padding:24px; line-height:1.65;
}
.container{max-width:1400px; margin:0 auto}
header{display:flex; align-items:baseline; gap:16px; flex-wrap:wrap; margin-bottom:20px}
h1{margin:0; font-size:22px}
.subtle{color:var(--muted); font-size:13px}
nav{margin:16px 0 24px; padding:12px 16px; background:var(--panel); border:1px solid var(--border); border-radius:8px}
nav a{color:var(--accent); margin-right:14px; text-decoration:none; font-size:13px; white-space:nowrap}
nav a:hover{text-decoration:underline}
.search-wrap{margin:0 0 12px}
input.search{width:100%; padding:10px 12px; background:var(--panel); color:var(--text);
  border:1px solid var(--border); border-radius:6px; font-size:14px}
.table-wrap{overflow-x:auto; background:var(--panel); border:1px solid var(--border); border-radius:8px}
table{border-collapse:collapse; width:100%; font-size:13px}
th,td{padding:10px 12px; border-bottom:1px solid var(--border); vertical-align:top; text-align:left}
th{background:var(--header-bg); color:var(--header-text); position:sticky; top:0; font-weight:600; cursor:pointer; user-select:none}
th:hover{background:#2a4f8c}
td{white-space:pre-wrap; word-break:break-word; max-width:480px}
tr:nth-child(even){background:rgba(127,127,127,0.04)}
tr:hover{background:rgba(88,166,255,0.08)}
a{color:var(--accent)}
.section{background:var(--panel); border:1px solid var(--border); border-radius:8px; padding:18px 22px; margin-bottom:14px}
.section h2{margin:0 0 8px; font-size:17px}
.section .meta{color:var(--muted); font-size:12px; margin-bottom:10px}
.section dl{margin:0; display:grid; grid-template-columns:max-content 1fr; gap:6px 16px; font-size:13px}
.section dt{color:var(--muted); white-space:nowrap}
.section dd{margin:0; white-space:pre-wrap}
.empty{color:var(--muted); font-style:italic; padding:30px; text-align:center}
footer{margin-top:30px; color:var(--muted); font-size:12px; text-align:center}
.tag{display:inline-block; background:rgba(88,166,255,0.15); color:var(--accent-2);
  padding:2px 8px; border-radius:10px; font-size:11px; margin-right:4px; margin-bottom:3px}
</style>
"""

JS = """
<script>
(function(){
  const inp=document.getElementById('search');
  if(!inp) return;
  inp.addEventListener('input',()=>{
    const q=inp.value.toLowerCase();
    document.querySelectorAll('.searchable').forEach(el=>{
      el.style.display = el.textContent.toLowerCase().includes(q) ? '' : 'none';
    });
  });
  // sort
  document.querySelectorAll('table').forEach(table=>{
    table.querySelectorAll('th').forEach((th,idx)=>{
      let asc=true;
      th.addEventListener('click',()=>{
        const tbody=table.tBodies[0];
        const rows=Array.from(tbody.rows);
        rows.sort((a,b)=>{
          const av=(a.cells[idx]?.innerText||'').trim();
          const bv=(b.cells[idx]?.innerText||'').trim();
          const an=parseFloat(av.replace(/,/g,''));
          const bn=parseFloat(bv.replace(/,/g,''));
          if(!isNaN(an)&&!isNaN(bn)) return (an-bn)*(asc?1:-1);
          return av.localeCompare(bv,'ja')*(asc?1:-1);
        });
        rows.forEach(r=>tbody.appendChild(r));
        asc=!asc;
      });
    });
  });
})();
</script>
"""

def nav_html(active_slug=None):
    items = []
    items.append(('00-readme', 'TOP'))
    for sn, slug in SHEET_SLUGS.items():
        if sn == "README":
            continue
        items.append((slug, sn))
    parts = []
    for slug, label in items:
        href = "index.html" if slug == "00-readme" else f"{slug}.html"
        parts.append(f'<a href="{href}">{html.escape(label)}</a>')
    return '<nav>' + ''.join(parts) + '</nav>'

def url_to_link(text):
    if not isinstance(text, str): return html.escape(str(text)) if text is not None else ''
    pattern = re.compile(r'(https?://[^\s<>"]+)')
    parts = []
    last = 0
    for m in pattern.finditer(text):
        parts.append(html.escape(text[last:m.start()]))
        url = m.group(1)
        parts.append(f'<a href="{html.escape(url)}" target="_blank" rel="noopener">{html.escape(url[:60] + ("…" if len(url)>60 else ""))}</a>')
        last = m.end()
    parts.append(html.escape(text[last:]))
    return ''.join(parts)

def fmt_cell(v):
    if v is None: return ''
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int):
        return f"{v:,}"
    if isinstance(v, float):
        return f"{v:,.0f}" if v == int(v) else f"{v:,}"
    return str(v)

def fmt_tag_cell(v):
    text = fmt_cell(v)
    if not text: return ''
    if any(sep in text for sep in [',', '、']):
        tags = re.split(r'[,、]\s*', text)
        return ''.join(f'<span class="tag">{html.escape(t.strip())}</span>' for t in tags if t.strip())
    return html.escape(text)

def render_table_html(sheet_name, headers, rows):
    nav = nav_html(SHEET_SLUGS.get(sheet_name))
    title = html.escape(sheet_name)
    if not rows:
        body = '<div class="empty">データなし</div>'
    else:
        ths = ''.join(f'<th>{html.escape(h or "")}</th>' for h in headers)
        trs = []
        for row in rows:
            tds = []
            for i, v in enumerate(row):
                h = headers[i] if i < len(headers) else ''
                text = fmt_cell(v)
                if not text:
                    cell_html = ''
                elif h and ('URL' in h or 'リンク' in h):
                    cell_html = url_to_link(text)
                elif h == 'タグ':
                    cell_html = fmt_tag_cell(v)
                else:
                    cell_html = url_to_link(text) if 'http' in text else html.escape(text)
                tds.append(f'<td>{cell_html}</td>')
            trs.append(f'<tr class="searchable">{"".join(tds)}</tr>')
        body = (
            '<div class="search-wrap"><input class="search" id="search" placeholder="全列を絞り込み検索…"></div>'
            '<div class="table-wrap"><table><thead><tr>'
            + ths + '</tr></thead><tbody>'
            + ''.join(trs) + '</tbody></table></div>'
        )
    updated = datetime.now().strftime('%Y-%m-%d %H:%M')
    return f"""<!DOCTYPE html><html lang="ja"><head><meta charset="utf-8">
<title>{title} | ゲーム業界トラッカー</title>
<meta name="viewport" content="width=device-width,initial-scale=1">
{CSS}</head><body><div class="container">
<header><h1>{title}</h1><span class="subtle">最終更新: {updated}</span></header>
{nav}{body}
<footer>ゲーム業界・時事情報収集 · 自動生成</footer>
</div>{JS}</body></html>"""

def render_section_html(sheet_name, headers, rows):
    nav = nav_html(SHEET_SLUGS.get(sheet_name))
    title = html.escape(sheet_name)
    sections = []
    for row in rows:
        d = {headers[i]: fmt_cell(v) for i, v in enumerate(row) if i < len(headers)}
        if not any(d.values()): continue
        head_field = headers[1] if len(headers) > 1 else headers[0]
        h2 = html.escape(d.get(head_field, '') or '(無題)')
        meta_parts = []
        if d.get('日付'): meta_parts.append(d['日付'])
        if d.get('記録日'): meta_parts.append(d['記録日'])
        if d.get('関連企業'): meta_parts.append(f"関連企業: {d['関連企業']}")
        if d.get('関連IP/タイトル') or d.get('関連企業/IP'):
            ip = d.get('関連IP/タイトル') or d.get('関連企業/IP')
            meta_parts.append(f"IP: {ip}")
        if d.get('タグ'): meta_parts.append(d['タグ'])
        if d.get('確度'): meta_parts.append(f"確度: {d['確度']}")
        meta = ' · '.join(meta_parts)
        dls = []
        skip_keys = {'日付', '記録日', '関連企業', '関連IP/タイトル', '関連企業/IP', 'タグ', '確度'}
        if head_field in d:
            skip_keys.add(head_field)
        for k in headers:
            if not k or k in skip_keys: continue
            v = d.get(k, '')
            if not v: continue
            if 'URL' in k or 'リンク' in k:
                vh = url_to_link(v)
            else:
                vh = url_to_link(v) if 'http' in v else html.escape(v).replace('\n','<br>')
            dls.append(f'<dt>{html.escape(k)}</dt><dd>{vh}</dd>')
        sections.append(f'<div class="section searchable"><h2>{h2}</h2>'
                        + (f'<div class="meta">{html.escape(meta)}</div>' if meta else '')
                        + (f'<dl>{"".join(dls)}</dl>' if dls else '') + '</div>')
    body = ('<div class="search-wrap"><input class="search" id="search" placeholder="全文検索…"></div>'
            + (''.join(sections) if sections else '<div class="empty">データなし</div>'))
    updated = datetime.now().strftime('%Y-%m-%d %H:%M')
    return f"""<!DOCTYPE html><html lang="ja"><head><meta charset="utf-8">
<title>{title} | ゲーム業界トラッカー</title>
<meta name="viewport" content="width=device-width,initial-scale=1">
{CSS}</head><body><div class="container">
<header><h1>{title}</h1><span class="subtle">最終更新: {updated}</span></header>
{nav}{body}
<footer>ゲーム業界・時事情報収集 · 自動生成</footer>
</div>{JS}</body></html>"""


RANKING_SHEET = "週間ランキング"
RANK_GROUP_ORDER = ["国内ファミ通", "Steamトップセラー", "NSW DL", "PS DL"]

def _week_key(s):
    m = re.search(r'(\d{4})-(\d{2})(?:-(\d{2}))?', str(s or ''))
    return m.group(0) if m else ''

def _rank_groups(headers, rows):
    idx = {h: i for i, h in enumerate(headers) if h}
    ki, wi, ri = idx.get('区分'), idx.get('集計週'), idx.get('順位')
    groups = {}
    for r in rows:
        g = fmt_cell(r[ki]) if ki is not None else 'その他'
        groups.setdefault(g or 'その他', []).append(r)
    order = [g for g in RANK_GROUP_ORDER if g in groups] + [g for g in groups if g not in RANK_GROUP_ORDER]
    def rank_val(r):
        try: return int(str(r[ri]).replace(',', ''))
        except Exception: return 999
    for g in groups:
        groups[g].sort(key=rank_val)
        groups[g].sort(key=lambda r: _week_key(fmt_cell(r[wi])), reverse=True)
    return idx, wi, order, groups

def render_ranking_html(sheet_name, headers, rows):
    nav = nav_html(SHEET_SLUGS.get(sheet_name))
    title = html.escape(sheet_name)
    idx, wi, order, groups = _rank_groups(headers, rows)
    disp = [h for h in headers if h and h not in ('区分', '集計週')]
    sections = []
    for gi, g in enumerate(order):
        grows = groups[g]
        weeks = sorted({fmt_cell(r[wi]) for r in grows}, key=_week_key, reverse=True)
        opts = ''.join(
            f'<option value="{html.escape(w)}"{" selected" if i == 0 else ""}>{html.escape(w)}</option>'
            for i, w in enumerate(weeks))
        opts += '<option value="__all__">すべての週</option>'
        ths = ''.join(f'<th>{html.escape(h)}</th>' for h in disp)
        trs = []
        for r in grows:
            week = fmt_cell(r[wi])
            tds = []
            for h in disp:
                v = r[idx[h]] if idx.get(h) is not None and idx[h] < len(r) else None
                text = fmt_cell(v)
                if not text:
                    cell = ''
                elif 'URL' in h or 'リンク' in h or '出典' in h:
                    cell = url_to_link(text)
                else:
                    cell = url_to_link(text) if 'http' in text else html.escape(text)
                tds.append(f'<td>{cell}</td>')
            trs.append(f'<tr data-group="g{gi}" data-week="{html.escape(week)}">{"".join(tds)}</tr>')
        sections.append(
            f'<div class="section"><h2>{html.escape(g)}</h2>'
            f'<div class="meta" style="margin-bottom:10px">対象週: '
            f'<select class="week-sel" data-group="g{gi}" '
            f'style="padding:6px 10px;background:var(--panel);color:var(--text);'
            f'border:1px solid var(--border);border-radius:6px;font-size:13px">{opts}</select></div>'
            f'<div class="table-wrap"><table><thead><tr>{ths}</tr></thead>'
            f'<tbody>{"".join(trs)}</tbody></table></div></div>')
    body = ''.join(sections) if sections else '<div class="empty">データなし</div>'
    week_js = """
<script>
document.querySelectorAll('select.week-sel').forEach(sel=>{
  const apply=()=>{
    const v=sel.value;
    document.querySelectorAll('tr[data-group=\"'+sel.dataset.group+'\"]').forEach(tr=>{
      tr.style.display=(v==='__all__'||tr.dataset.week===v)?'':'none';
    });
  };
  sel.addEventListener('change',apply); apply();
});
</script>"""
    updated = datetime.now().strftime('%Y-%m-%d %H:%M')
    return f"""<!DOCTYPE html><html lang="ja"><head><meta charset="utf-8">
<title>{title} | ゲーム業界トラッカー</title>
<meta name="viewport" content="width=device-width,initial-scale=1">
{CSS}</head><body><div class="container">
<header><h1>{title}</h1><span class="subtle">最終更新: {updated}</span></header>
{nav}{body}
<footer>ゲーム業界・時事情報収集 · 自動生成</footer>
</div>{JS}{week_js}</body></html>"""

def render_ranking_md(sheet_name, headers, rows):
    lines = [f"# {sheet_name}", "", f"_最終更新: {datetime.now().strftime('%Y-%m-%d %H:%M')}_", ""]
    idx, wi, order, groups = _rank_groups(headers, rows)
    disp = [h for h in headers if h and h not in ('区分', '集計週')]
    for g in order:
        lines.append(f"## {g}")
        lines.append("")
        weeks = sorted({fmt_cell(r[wi]) for r in groups[g]}, key=_week_key, reverse=True)
        for w in weeks:
            lines.append(f"### {w}")
            lines.append("")
            lines.append('| ' + ' | '.join(disp) + ' |')
            lines.append('| ' + ' | '.join(['---'] * len(disp)) + ' |')
            for r in groups[g]:
                if fmt_cell(r[wi]) != w: continue
                cells = []
                for h in disp:
                    v = r[idx[h]] if idx.get(h) is not None and idx[h] < len(r) else None
                    cells.append(fmt_cell(v).replace('|', '\\|').replace('\n', '<br>'))
                lines.append('| ' + ' | '.join(cells) + ' |')
            lines.append("")
    return '\n'.join(lines) + '\n'

def render_table_md(sheet_name, headers, rows):
    lines = [f"# {sheet_name}", "", f"_最終更新: {datetime.now().strftime('%Y-%m-%d %H:%M')}_", ""]
    safe_headers = [h or '' for h in headers]
    lines.append('| ' + ' | '.join(safe_headers) + ' |')
    lines.append('| ' + ' | '.join(['---'] * len(safe_headers)) + ' |')
    for row in rows:
        cells = []
        for v in row:
            t = fmt_cell(v).replace('|', '\\|').replace('\n', '<br>')
            cells.append(t)
        lines.append('| ' + ' | '.join(cells) + ' |')
    return '\n'.join(lines) + '\n'

def render_section_md(sheet_name, headers, rows):
    lines = [f"# {sheet_name}", "", f"_最終更新: {datetime.now().strftime('%Y-%m-%d %H:%M')}_", ""]
    for row in rows:
        d = {headers[i]: fmt_cell(v) for i, v in enumerate(row) if i < len(headers)}
        if not any(d.values()): continue
        head_field = headers[1] if len(headers) > 1 else headers[0]
        lines.append(f"## {d.get(head_field, '(無題)')}")
        for k in headers:
            if not k or k == head_field: continue
            v = d.get(k, '')
            if not v: continue
            lines.append(f"- **{k}**: {v}")
        lines.append("")
    return '\n'.join(lines) + '\n'

def render_index_html(stats):
    nav = nav_html('00-readme')
    rows = []
    for sn, slug in SHEET_SLUGS.items():
        if sn == 'README': continue
        n = stats.get(sn, 0)
        rows.append(f'<tr><td><a href="{slug}.html">{html.escape(sn)}</a></td>'
                    f'<td style="text-align:right">{n}</td></tr>')
    updated = datetime.now().strftime('%Y-%m-%d %H:%M')
    return f"""<!DOCTYPE html><html lang="ja"><head><meta charset="utf-8">
<title>ゲーム業界・時事情報収集</title>
<meta name="viewport" content="width=device-width,initial-scale=1">
{CSS}</head><body><div class="container">
<header><h1>🎮 ゲーム業界・時事情報収集</h1><span class="subtle">最終更新: {updated}</span></header>
{nav}
<div class="section">
  <h2>シート一覧</h2>
  <div class="table-wrap"><table><thead><tr><th>シート</th><th style="text-align:right">行数</th></tr></thead>
  <tbody>{''.join(rows)}</tbody></table></div>
</div>
<div class="section">
  <h2>更新スケジュール</h2>
  <ul>
    <li><b>毎朝 7:00</b> — 日次ニュース10件追加</li>
    <li><b>毎週月 9:00</b> — 国内/海外大手の株価・IR更新</li>
    <li><b>毎週金 18:00</b> — 新作リリース週次レポート</li>
    <li><b>毎週水 12:00</b> — 業界トレンド・話題のゲームデザイン</li>
    <li><b>2/5/8/11月 上旬</b> — 四半期決算サマリ</li>
  </ul>
</div>
<footer>自動生成 · GitHub Pages</footer>
</div></body></html>"""

def render_index_md(stats):
    lines = ["# 🎮 ゲーム業界・時事情報収集", "",
             f"_最終更新: {datetime.now().strftime('%Y-%m-%d %H:%M')}_", "",
             "## シート一覧", ""]
    for sn, slug in SHEET_SLUGS.items():
        if sn == 'README': continue
        n = stats.get(sn, 0)
        lines.append(f"- [{sn}](md/{slug}.md) ({n}行)")
    lines += ["", "## 更新スケジュール", "",
              "- **毎朝 7:00** — 日次ニュース10件追加",
              "- **毎週月 9:00** — 国内/海外大手の株価・IR更新",
              "- **毎週金 18:00** — 新作リリース週次レポート",
              "- **毎週水 12:00** — 業界トレンド・話題のゲームデザイン",
              "- **2/5/8/11月 上旬** — 四半期決算サマリ", "",
              "ダッシュボード（GitHub Pages）から閲覧できます。"]
    return '\n'.join(lines) + '\n'

def main():
    if not XLSX_PATH.exists():
        print(f"ERROR: not found: {XLSX_PATH}", file=sys.stderr); sys.exit(1)
    MD_DIR.mkdir(parents=True, exist_ok=True)
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(XLSX_PATH, data_only=True)
    stats = {}
    for sn in wb.sheetnames:
        if sn not in SHEET_SLUGS: continue
        if sn == "README": continue
        ws = wb[sn]
        headers = [c.value for c in ws[1]]
        # 末尾の全Noneヘッダはトリム
        while headers and headers[-1] is None:
            headers.pop()
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            row = list(row[:len(headers)])
            if any(v not in (None, '') for v in row):
                rows.append(row)
        # 全シートを日付列を基準に新しい順（降順）で並べ替え（最新を上に）。
        # 各シートで最初に見つかった日付系ヘッダを基準にする。
        # YYYY-MM-DD形式を抽出してキーに使用。日付として解釈できない行は最下段。
        # 週間ランキングは専用レンダラ側でソートするため除外。
        DATE_HEADER_PRIORITY = ['日付', '記録日', '発売日', '更新日', '決算発表日',
                                '開催日', '集計週', '最終更新日']
        if sn != RANKING_SHEET:
            date_col = None
            for _cand in DATE_HEADER_PRIORITY:
                if _cand in headers:
                    date_col = headers.index(_cand)
                    break
            if date_col is not None:
                _date_pat = re.compile(r'(\d{4}-\d{2}-\d{2})')
                def _date_key(r, _c=date_col):
                    v = r[_c] if _c < len(r) else None
                    if isinstance(v, datetime):
                        return v.strftime('%Y-%m-%d')
                    s = str(v) if v is not None else ''
                    m = _date_pat.search(s)
                    return m.group(1) if m else ''
                rows.sort(key=_date_key, reverse=True)
        stats[sn] = len(rows)
        slug = SHEET_SLUGS[sn]
        if sn == RANKING_SHEET:
            (MD_DIR / f"{slug}.md").write_text(render_ranking_md(sn, headers, rows), encoding='utf-8')
            (DOCS_DIR / f"{slug}.html").write_text(render_ranking_html(sn, headers, rows), encoding='utf-8')
        elif sn in SECTION_SHEETS:
            (MD_DIR / f"{slug}.md").write_text(render_section_md(sn, headers, rows), encoding='utf-8')
            (DOCS_DIR / f"{slug}.html").write_text(render_section_html(sn, headers, rows), encoding='utf-8')
        else:
            (MD_DIR / f"{slug}.md").write_text(render_table_md(sn, headers, rows), encoding='utf-8')
            (DOCS_DIR / f"{slug}.html").write_text(render_table_html(sn, headers, rows), encoding='utf-8')
    (DOCS_DIR / "index.html").write_text(render_index_html(stats), encoding='utf-8')
    (REPO_ROOT / "README.md").write_text(render_index_md(stats), encoding='utf-8')
    print(f"Generated for {len(stats)} sheets. Total rows: {sum(stats.values())}")

if __name__ == "__main__":
    main()
