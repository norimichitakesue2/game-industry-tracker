# -*- coding: utf-8 -*-
import openpyxl

PATH = "master/ゲーム業界・時事情報収集.xlsx"
wb = openpyxl.load_workbook(PATH)

def norm(x): return "" if x is None else str(x).strip()

def append_dedup(sheet, rows, key_idx):
    ws = wb[sheet]
    existing = set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        existing.add(tuple(norm(r[i]) for i in key_idx))
    added = 0
    for row in rows:
        k = tuple(norm(row[i]) for i in key_idx)
        if k in existing: continue
        ws.append(list(row)); existing.add(k); added += 1
    print(f"[{sheet}] appended {added}/{len(rows)}")
    return added

def update_rows(sheet, match_idx, match_val, updates):
    ws = wb[sheet]; hit = 0
    for r in ws.iter_rows(min_row=2):
        if norm(r[match_idx].value) == match_val:
            for ci, val in updates.items(): r[ci].value = val
            hit += 1
    print(f"[{sheet}] updated {hit} row(s) matching '{match_val}'")
    return hit

summary = {}

update_rows("新作リリースログ", 1, "Silent Hill: Townfall（サイレントヒル タウンフォール）", {
    6: 81, 7: 81,
    11: "サイコロジカルホラー。Metacritic/OpenCritic ともに81(OC 86パーセンタイル・推奨83%)。雰囲気・サウンド・心理的恐怖を高評価も、ステルス戦とチェックポイントの渋さに批判。ファミ通週販4位(初週13,596本)。Digital Spy等は満点。",
    9: "初週13,596本(ファミ通/国内PS5)"})
update_rows("新作リリースログ", 1, "ACE COMBAT 8: WINGS OF THEVE（エースコンバット8 ウイングス・オブ・シーヴ）", {
    0: "2026-10-02", 6: 87, 7: 88,
    11: "シリーズナンバリング最新作。Metacritic87(PS5)/OpenCritic88・推奨98%で“25年ぶりのシリーズ最高級”(AC04の89点に次ぐ)。空戦・楽曲・ストーリーを絶賛、演出のテンポに一部指摘。Steam週間トップセラー8位(予約)。"})

rel_rows = [
 ("2026-10-05(予定)","AION 2（アイオン2）","NC","NC","PC/モバイル","MMORPG(基本無料)",None,None,None,None,3,
  "【来週】NCSOFTの人気MMO続編。基本無料。Steam週間トップセラーに予約でNEWランクイン。※AI生成コンテンツ開示あり。","https://store.steampowered.com/app/3393110/AION_2/"),
 ("2026-10-08(予定)","KINGDOM HEARTS Collection [I～III]（キングダム ハーツ コレクション）","スクウェア・エニックス","スクウェア・エニックス","PS5/Switch2/PC/Xbox","アクションRPG(コレクション)",None,None,None,None,4,
  "【来週の大型】KHシリーズ主要作を収録したコレクション。10/8発売予定。","https://www.gamegrin.com/directory/game/kingdom-hearts-collection-iiii/about"),
 ("2026-10-09(予定)","Dragon's Dogma 2: Dark Arisen（ドラゴンズドグマ2：ダークアリズン）","カプコン","カプコン","PS5/Xbox/PC","アクションRPG(拡張)",None,None,None,None,4,
  "【来週の大型】『ドラゴンズドグマ2』の大型拡張/強化版。10/9発売予定。","https://www.gamegrin.com/directory/game/dragons-dogma-2-dark-arisen/about"),
]
summary["新作リリースログ"] = append_dedup("新作リリースログ", rel_rows, (0,1))

indie_rows = [
 ("2026-10-02","Dressmaker（ドレスメーカー）","Cozy Lives","2026-09-21 配信","PC(Steam)","カジュアル/クラフト・ショップ経営シム",
  "配信直後にSteam週間トップセラー上位(NEW)入り。『Broforce』で知られるFree Livesがパブリッシュする仕立屋ショップ運営のコージーゲーム。",
  "Steam週間トップセラーNEWランクイン(2026-09-22〜09-29)","Free Lives(パブ)×Cozy Lives。コージー系の話題作。継続的な同接・レビューを要ウォッチ。"),
 ("2026-10-02","Graveyard Keeper 2（グレイブヤード キーパー2）","Lazy Bear Games","2026-09-22 配信","PC(Steam)","サンドボックス/経営シム",
  "墓守経営シムの続編。配信直後にSteam週間トップセラー上位(NEW)入り。パブはtinyBuild。",
  "Steam週間トップセラーNEWランクイン(2026-09-22〜09-29)","前作が人気のインディー続編。製品版評価が固まれば要注目。"),
]
summary["インディー注目"] = append_dedup("インディー注目", indie_rows, (1,))

FAMI_SRC="https://www.famitsu.com/article/202610/89721"
STEAM_SRC="https://www.gamegrin.com/news/weekly-top-selling-games-on-steam-22nd29th-of-september-2026"
fami_week="2026-09-21〜09-27"
fami=[(1,"ダービースタリオン2","Switch2",29336,29336),(2,"リズム天国 ミラクルスターズ","Switch",27309,1049922),
 (3,"ファイアーエムブレム 万紫千紅","Switch2",26728,153859),(4,"SILENT HILL: Townfall","PS5",13596,13596),
 (5,"ドラゴンクエストXI 過ぎ去りし時を求めて S","Switch2",10844,10844),(6,"EA SPORTS FC 27","PS5",9873,9873),
 (7,"トモダチコレクション わくわく生活","Switch",8052,1618104),(8,"スプラトゥーン レイダース","Switch2",7405,671202),
 (9,"EA SPORTS FC 27","Switch2",5952,5952),(10,"鬼武者 Way of the Sword","PS5",5141,110087)]
rank_rows=[]
for rk,t,p,wk,cum in fami:
    rank_rows.append((fami_week,"国内ファミ通",rk,t,p,wk,cum,None,FAMI_SRC,"ファミ通週間ソフト推定販売本数TOP10(初登場5本/4作品)。首位は競走馬育成シム最新作『ダービースタリオン2』。"))
steam_week="2026-09-22〜09-29"
steam=[(1,"EA SPORTS FC 27"),(2,"Counter-Strike 2"),(3,"WARDOGS"),(4,"CONTROL Resonant"),(5,"Apex Legends"),
 (6,"Total War: WARHAMMER III"),(7,"Aniimo"),(8,"ACE COMBAT 8: WINGS OF THEVE"),(9,"Warframe"),(10,"PUBG: BATTLEGROUNDS")]
for rk,t in steam:
    note="Steam週間トップセラー(GameGrin集計、収益ベース、順位のみ)。EA FC27が1位デビュー。"
    if t=="ACE COMBAT 8: WINGS OF THEVE": note="Steam週間トップセラー(GameGrin集計)。本作は予約段階でのランクイン。"
    rank_rows.append((steam_week,"Steamトップセラー",rk,t,"PC(Steam)",None,None,None,STEAM_SRC,note))
summary["週間ランキング"]=append_dedup("週間ランキング",rank_rows,(0,1,2))

EIGA="https://eiga.com/news/20260928/28/"; CT="https://www.cinematoday.jp/news/N0157107"
movie_rows=[
 ("2026-09-26","ハート・オブ・ビースト",None,None,None,None,"週末興行 初登場4位",None,None,None,3,"9/25-27の週末、国内興行で初登場4位。",EIGA),
 ("2026-09-26","マッチング TRUE LOVE",None,None,"日本",None,"週末興行 初登場7位",None,None,None,2,"9/25-27の週末、国内興行で初登場7位。",EIGA),
 ("2026-09-26","アベンジャーズ エンドゲーム アンコール（Avengers: Endgame Encore）",None,"ディズニー","米国","アクション(再上映)","週末興行 初登場5位",None,None,None,3,"MCU『エンドゲーム』のアンコール上映。9/25-27週末、国内興行で初登場5位。",CT),
]
summary["【映画】公開・興行"]=append_dedup("【映画】公開・興行",movie_rows,(0,1))

anime_rows=[
 ("2026-10-01","FX戦士くるみちゃん","大童澄瞳ほか/Web発",None,"2026年秋(10月期)","各種配信","ギャグ/青春",None,None,3,
  "【秋アニメ新番組】FXで2000万円を溶かした女子大生が“戦場”へ向かう異色のマネーギャグ。10/1放送開始。","https://animeanime.jp/article/2026/10/01/103411.html"),
 ("2026-10-02","蒼き伝承 ‒ Welsh & Shedar（蒼き伝承）","Ankama(オリジナル)","Ankama","2026年秋(10月期)","TOKYO MXほか","冒険/ファンタジー",None,None,3,
  "【秋アニメ新番組】仏Ankamaが贈る冒険ファンタジー。10/2 TOKYO MXで放送開始。","https://animeanime.jp/article/2026/10/02/103442.html"),
]
summary["【アニメ】放送・配信ログ"]=append_dedup("【アニメ】放送・配信ログ",anime_rows,(0,1))

ORICON_C="https://www.oricon.co.jp/rank/obc/w/2026-09-28/"
manga_rows=[
 ("2026-09","The JOJOLands 9","荒木飛呂彦","集英社","ウルトラジャンプ","9",None,"オリコン週間コミック(単行本)1位(2026-09-28付/集計9/14-20)・推定54,642部",4,"ジョジョ第9部最新巻が単行本ランキング首位。",ORICON_C),
 ("2026-09","女の園の星 5","和山やま","祥伝社","FEEL YOUNG","5",None,"オリコン週間コミック(単行本)2位(2026-09-28付/集計9/14-20)・推定39,144部",3,"人気学園コメディ最新巻。単行本ランキング2位。",ORICON_C),
 ("2026-09","うるわしの宵の月 11","やまもり三香","講談社","デザート","11",None,"オリコン週間コミック(単行本)3位(2026-09-28付/集計9/14-20)・推定36,210部",3,"少女漫画の人気作。単行本ランキング3位。",ORICON_C),
]
summary["【漫画】新刊・話題作"]=append_dedup("【漫画】新刊・話題作",manga_rows,(0,1))

novel_rows=[
 ("2026-09-09","白い夜のセレナーデ","赤川次郎","光文社","光文社文庫","新刊",None,None,2,"35年以上毎年刊行される赤川次郎の長寿シリーズ第39作。2026年9月の文庫新刊。","https://hon-hikidashi.jp/news/177001/"),
]
summary["【小説】新刊・受賞"]=append_dedup("【小説】新刊・受賞",novel_rows,(0,1))

BJ="https://www.billboard-japan.com/charts/detail?a=hot100"; BB="https://www.billboard.com/charts/hot-100/"
music_rows=[
 ("2026-10-05(付)","チャート","REBON","Number_i","Atlantic Records","シングル/配信","Billboard JAPAN Hot 100 1位(最新週)",None,5,"Atlantic Records移籍第一弾。『REBON』『DIGITAL GIRL』で総合チャート1-2位を独占。",BJ),
 ("2026-10-05(付)","チャート","DIGITAL GIRL","Number_i","Atlantic Records","シングル/配信","Billboard JAPAN Hot 100 2位(最新週)",None,4,"Number_iが1位(REBON)と合わせ総合チャート上位を独占。",BJ),
 ("2026-10(フラゲ)","新譜","POPS","Mrs. GREEN APPLE","ユニバーサル","アルバム","先ヨミでアルバム首位独走(60.2万枚)",None,5,"フラゲ日集計で58万枚突破、2作目のハーフミリオン達成見込みの大型新譜。","https://www.billboard-japan.com/d_news/detail/166068/"),
 ("2026-10-03(付)","チャート","Choosin' Texas","Ella Langley",None,"シングル/配信","Billboard Hot 100(米) 1位・通算24週目",None,4,"全米Billboard Hot 100で通算24週目の首位というロングラン。",BB),
]
summary["【音楽】リリース・チャート"]=append_dedup("【音楽】リリース・チャート",music_rows,(2,3))

wb.save(PATH)
print("\n=== SUMMARY ===")
for k,v in summary.items(): print(f"  {k}: +{v}")
