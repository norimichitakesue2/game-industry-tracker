# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, Alignment
from datetime import date
TODAY="2026-08-18"
PATH="master/ゲーム業界・時事情報収集.xlsx"
wb=openpyxl.load_workbook(PATH)
FONT=Font(name="Yu Gothic", size=11)
ALIGN=Alignment(wrap_text=True, vertical="top")
def style_row(ws,r):
    for c in range(1,ws.max_column+1):
        ws.cell(r,c).font=FONT
        ws.cell(r,c).alignment=ALIGN

log=[]

# ===== 横断的問い =====
ws=wb["横断的問い"]
existing_q=[str(ws.cell(r,2).value or "") for r in range(2,ws.max_row+1)]
def q_dup(q):
    key=q[:20]
    return any(key in e for e in existing_q)
QS=[
[TODAY,
 "主要音楽DSP（Spotify/Amazon/YouTube/Tidal）の一斉値上げでサブスク成長が“値上げ依存”に転じた構図は、ゲームのサブスク（Game Pass/PS Plus）や動画配信でも同じ加入者純増の頭打ち→単価依存に陥るか？",
 "UMG株25%急落・成長の値上げ依存（財経新聞8/1）／Tidal値上げ8/3",
 "Universal Music,Spotify,Tidal,Microsoft,Sony",
 "通信・SaaSの価格改定依存／サブスク疲れ",
 "未検証","エンタメ各サブスクが同時に“純増限界→値上げ”局面に入りつつある可能性。ゲーム/動画のARPU戦略を並べて観測したい","横断"],
[TODAY,
 "UMGが導入する新譜72時間の有料会員先行ウィンドウ（無料層への時間差開放）は、映画の劇場→配信ウィンドウやゲームの早期アクセス／デラックス版先行と同じ“窓口戦略”の音楽版か？各エンタメの最適ウィンドウはどこに収束するか？",
 "UMGの72時間先行ウィンドウ導入（Chartlex）",
 "Universal Music Group,Spotify",
 "映画の公開ウィンドウ／ゲームの早期アクセス／出版の先行配信",
 "未検証","同一コンテンツを時間差・階層で出し分ける設計が全エンタメで共通言語化しつつある","横断"],
[TODAY,
 "Crunchyrollの無告知デリスト（配信棚からの作品消失）は、音楽サブスクの楽曲消失やゲームのストア配信終了と同じ“所有なき視聴”の不安定性か？ユーザーの所有欲（物理・買い切り）回帰をどのジャンルで最も強く促すか？",
 "『86』Crunchyroll復活＋7作品を無告知削除（ANN 8/17）",
 "Crunchyroll,Sony,Spotify",
 "小売のサブスク疲れ→所有回帰／中古・コレクター市場",
 "未検証","サイレントメビウスの高単価BD化など“物理プレミアム”回帰と同じ根を持つ可能性","横断"],
]
addq=0
for row in QS:
    if q_dup(row[1]): continue
    ws.append(row); style_row(ws,ws.max_row); addq+=1
log.append(("横断的問い",addq))

# ===== メディアミックス追跡 =====
ws=wb["メディアミックス追跡"]
existing_mm=set()
for r in range(2,ws.max_row+1):
    existing_mm.add((str(ws.cell(r,2).value or "").strip(), str(ws.cell(r,5).value or "").strip()))
MM=[
[TODAY,"キングダム ハーツ","ゲーム","スクウェア・エニックス／ディズニー","アニメ","シリーズ初のアニメーションシリーズ制作決定（KH4と同時発表、Disney+/Disney Channel想定）","2026-08-16","未定","発表","スクウェア・エニックス,ディズニー","KH4は2027年後半発売。ゲームIPの自社主導アニメ化"],
[TODAY,"ヒマチの嬢王","漫画","くれちはる／小学館（マンガワン）","実写ドラマ","永野芽郁 主演＆初プロデュースのDMM TVオリジナルドラマ","2026-08-17","2026年冬","発表","DMM TV,小学館","累計120万部超・1.7億PVのアプリ発ヒット。俳優が企画から参画"],
[TODAY,"赤ずきん、旅の途中で死体と出会う。","小説","青柳碧人","アニメ","2027年TVアニメ化（制作BENTEN Film、宮本侑芽・関根明良）","2026-08-10","2027年","発表","BENTEN Film,KADOKAWA","童話×本格ミステリ。実写化も並行進行の多面展開"],
]
addmm=0
for row in MM:
    key=(row[1].strip(),row[4].strip())
    if key in existing_mm: continue
    ws.append(row); style_row(ws,ws.max_row); existing_mm.add(key); addmm+=1
log.append(("メディアミックス追跡",addmm))

# ===== 今後起きそうなこと =====
ws=wb["今後起きそうなこと"]
existing_p=[str(ws.cell(r,2).value or "") for r in range(2,ws.max_row+1)]
def p_dup(t):
    key=t[:22]
    return any(key in e for e in existing_p)
PRED=[
[TODAY,
 "主要DSPの一斉値上げでサブスク純増が鈍り“価格改定依存”が露呈（UMG時価総額25%減）したことを起点に、DSP各社は解約抑止のため通信・デバイスとのバンドルや独占ウィンドウ確保に走り、音楽流通の主導権がレーベル／DSPからキャリア・ハード事業者側へ一部移る",
 "中","2026-2027",
 "Universal Music,Spotify,Tidal,Apple,通信キャリア",
 "「UMG株25%急落・成長の値上げ依存」（財経新聞8/1）／「Tidal値上げ8/3」（What Hi-Fi）",
 "各DSPの新規バンドル提携発表・ARPU/解約率の開示、キャリアの音楽サブスク囲い込み施策の有無を観測。純増鈍化がバンドル競争に転化するかを見る",
 "音楽"],
]
addp=0
for row in PRED:
    if p_dup(row[1]): continue
    ws.append(row); style_row(ws,ws.max_row); addp+=1
log.append(("今後起きそうなこと",addp))

# ===== イベントカレンダー =====
ws=wb["イベントカレンダー"]
existing_ev=set(str(ws.cell(r,3).value or "").strip() for r in range(2,ws.max_row+1))
NEWEV=[
["2026-08-18","2026-08-18","PlayStation State of Play (Phantom Blade Zero)","ショーケース","オンライン","配信","『Phantom Blade Zero』実機20分。予約開始・10/29発売決定","https://blog.playstation.com/2026/08/11/watch-the-phantom-blade-zero-gameplay-deep-dive-state-of-play-on-august-17-pre-orders-live-today/","終了",TODAY],
["2026-08-14","2026-08-17","Busan Indie Connect Festival 2026 (BIC2026)","展示会","釜山(韓国)","現地","応募総数が過去最大を更新。41の国と地域から180作品以上が展示","https://www.4gamer.net/","終了",TODAY],
]
adde=0
for row in NEWEV:
    if row[2].strip() in existing_ev: continue
    ws.append(row); style_row(ws,ws.max_row); existing_ev.add(row[2].strip()); adde+=1

# status maintenance for existing rows
def parse(d):
    d=str(d or "").strip()
    if len(d)==10 and d[4]=='-' and d[7]=='-':
        try: return date(int(d[:4]),int(d[5:7]),int(d[8:10]))
        except: return None
    return None
t=parse(TODAY)
upd=0
for r in range(2,ws.max_row+1):
    sd=parse(ws.cell(r,1).value); ed=parse(ws.cell(r,2).value)
    cur=str(ws.cell(r,9).value or "").strip()
    if cur in ("日程未確定","中止","延期"): continue
    if sd is None: continue
    end=ed or sd
    if end < t: new="終了"
    elif sd <= t <= end: new="開催中"
    elif sd > t: new="開催予定"
    else: new=cur
    if new and new!=cur:
        ws.cell(r,9).value=new
        ws.cell(r,10).value=TODAY
        ws.cell(r,9).font=FONT; ws.cell(r,9).alignment=ALIGN
        ws.cell(r,10).font=FONT; ws.cell(r,10).alignment=ALIGN
        upd+=1
log.append(("イベント新規",adde))
log.append(("イベント状態更新",upd))

wb.save(PATH)
for k,v in log: print(k,v)
print("SAVED sheets2")
