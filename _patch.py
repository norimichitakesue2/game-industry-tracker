# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, Alignment
from copy import copy
TODAY="2026-08-04"; NP="\n"
PATH="master/ゲーム業界・時事情報収集.xlsx"
wb=openpyxl.load_workbook(PATH)
FONT=Font(name="Yu Gothic",size=10); ALIGN=Alignment(wrap_text=True,vertical="top")
ws=wb["日次ニュース"]
existing=set(str(r[3]).strip() for r in ws.iter_rows(min_row=2,values_only=True) if r[3])
rows=[
("アニメ","『機動警察パトレイバー』新シリーズが8月14日に劇場公開、人気IPを再起動",
 "https://animatetimes.com/tag/details.php?id=4105",
 "『機動警察パトレイバー』シリーズの新たなシリーズが8月14日に劇場公開予定。1980年代末から続くロボット×警察の名作IPを再起動し、往年のファンと新規層の双方を狙う劇場アニメとして展開される。",
 "休眠気味だったレガシーIPの劇場再起動は認知資産を活かしつつ配信・グッズへ広げる王道戦略。新規オリジナルの当たり外れを避け、既知IPの再現×刷新で確実な初動を狙う近年の潮流を映す事例といえる。",
 "8月14日に新シリーズを劇場公開"+NP+"パトレイバーの久々のシリーズ再起動"+NP+"レガシーIPの認知資産を再活用"+NP+"往年ファンと新規層の双方を狙う",
 "製作委員会は眠っていた名作IPの劇場再起動を宣伝効率の高い安全資産として続々投入する。新作オリジナルより既知IPの続編・リブートに資源が集中し、レガシーIPの棚卸しと権利再交渉が活発化する。",
 "レガシーロボIPの再起動と劇場アニメ採算","-","機動警察パトレイバー","新作, 国内"),
("小説","オリコン上半期“本”ランキング2026、文庫1位は夕木春央『方舟』が席巻",
 "https://www.oricon.co.jp/special/75283/",
 "オリコンの第19回上半期本ランキング2026が発表。BOOK総合1位は堀田秀吾の習慣本、文庫ランキング1位は夕木春央『方舟』となった。全国書店4092店＋WEB通販の売上を集計した半期の消費者動向を示す指標だ。",
 "実用書が総合首位、ミステリー文庫がロングセラー化する構図が示唆的。話題化→文庫化→SNS再燃の循環でヒットが長期化し、単行本より回転の速い文庫・実用が販売の柱に。原作在庫の映像化・再ヒットの土壌にもなる。",
 "上半期BOOK総合1位は習慣本（実用書）"+NP+"文庫1位は夕木春央『方舟』"+NP+"全国4092書店＋WEB通販で集計"+NP+"文庫・実用のロングセラー化が鮮明",
 "出版社は話題作の文庫化・実用書の反復ヒットを収益の回転軸に据える。SNS発の再燃やメディアミックスで旧作在庫が再評価され、ミステリー等の映像化・コミカライズ企画の供給源として文庫ヒットの重要度が増す。",
 "半期ベストセラーと文庫ロングセラー化","オリコン","方舟","チャート・ランキング, 出版, 国内"),
("小説","BOOK WALKER 2026年上半期ラノベ・新文芸ランキングTOP50発表、電子が主戦場に",
 "https://bookwalker.jp/ex/feature/year_ranking/genre-lanove.html",
 "電子書籍ストアBOOK WALKERが2026年上半期のライトノベル・新文芸ランキングTOP50を発表。なろう系・新文芸を中心に電子先行・電子主戦場化が進み、紙より電子での初速がヒットの起点になる構造が定着している。",
 "ラノベ/新文芸の購買が電子に大きく移り、電子ランキングがアニメ化候補の先行指標になっている点が示唆的。電子の初速と継続DLが出版社のアニメ化・コミカライズ判断を左右し、IP発掘の入口が電子に移っている。",
 "上半期ラノベ・新文芸TOP50を発表"+NP+"なろう系・新文芸が上位を占有"+NP+"電子先行・電子主戦場化が進行"+NP+"電子ランキングがアニメ化の先行指標に",
 "電子ラノベの初速データがメディアミックス投資の判断材料として重みを増す。出版社は電子ランキング上位作を優先的にアニメ化・コミカライズし、Webtoon/なろう発IPのアニメ化パイプラインが一段と太くなる。",
 "電子ラノベ市場とアニメ化先行指標","BOOK WALKER, KADOKAWA","-","チャート・ランキング, 出版, 国内"),
]
added={}
for genre,h,u,bg,s,p,f,st,c,ip,t in rows:
    if u.strip() in existing:
        print("skip dup",u); continue
    r=ws.max_row+1
    for ci,v in enumerate([TODAY,genre,h,u,bg,s,p,f,st,c,ip,t,""],start=1):
        cell=ws.cell(row=r,column=ci,value=v); cell.font=copy(FONT); cell.alignment=copy(ALIGN)
    added[genre]=added.get(genre,0)+1; existing.add(u.strip())
wb.save(PATH)
print("patch added:",added)
