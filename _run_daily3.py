# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, Alignment
TODAY="2026-09-14"; PATH="master/ゲーム業界・時事情報収集.xlsx"
wb=openpyxl.load_workbook(PATH)
FONT=Font(name="Yu Gothic", size=11); ALIGN=Alignment(wrap_text=True, vertical="top")
def sr(ws,r,n):
    for c in range(1,n+1):
        cell=ws.cell(row=r,column=c); cell.font=FONT; cell.alignment=ALIGN
NL="\n"
ws=wb["日次ニュース"]; existing=set()
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[3]: existing.add(str(row[3]).strip())
news=[
 [TODAY,"ゲーム","Steam注目インディー15選、北欧神話のヴァンサバ系ローグライトACTが正式リリース","https://news.yahoo.co.jp/articles/25ba446445b14bd67edeb305661799a5523d410e","Game*Sparkの『本日のSteam注目ゲーム』で、北欧神話がテーマのヴァンパイアサバイバーズ系ローグライトACTが正式リリースされるなど、インディー新作が日々多数登場している状況が示された。","『ヴァンサバライク』『好評率100%のローグライク』など、確立された型を踏襲した高完成度インディーが量産される段階。低コストで一定の遊びを担保する型が個人・小規模開発の主流になっている。","北欧神話ヴァンサバ系ACTが正式リリース"+NL+"好評率の高いローグライク系が続々"+NL+"確立ジャンルの量産インディーが活況","確立された型(ヴァンサバ/ローグライト)の量産で、インディーは差別化を演出・世界観・コラボに求める段階へ。埋没を避けるIP・話題性の設計が競争軸になる。","インディーのジャンル定型化と差別化","","","インディー,トレンド,海外","二次ソース(Steam新作紹介)"],
 [TODAY,"漫画","『週刊少年ジャンプ』で新連載が続々投入、巻頭カラーで新人作家を大量起用","https://shonenjumpplus.com/magazine/17107094913312840225","『週刊少年ジャンプ』では41号で新連載『グレイキースの魔界語訳録』、40号で『ウラのレポート』が巻頭カラー54ページで掲載されるなど、秋の新連載ラッシュで新人作家が相次ぎ投入されている。","看板誌が短いスパンで新連載を大量投入し、巻頭カラーで新人を送り出す新陳代謝の仕組み。ヒットIPの供給源として、連載枠の入れ替えと新人育成を高速で回している。","41号『グレイキースの魔界語訳録』新連載"+NL+"40号『ウラのレポート』新連載"+NL+"巻頭カラーで新人を大量起用","雑誌はIP創出のR&D装置として新連載を高速で回し、生き残った作品をアニメ・ゲーム・海外へ多面展開。新人育成と連載枠の新陳代謝が出版社の競争力の源泉になる。","看板誌の新連載新陳代謝とIP創出","集英社","週刊少年ジャンプ","出版,新作,国内",""],
 [TODAY,"小説","村上春樹が翻訳を手がけた絵本『あらゆるものがおくりもの』が9/17発売","https://www.j-n.co.jp/news/754/","パトリック・マクドネル作の絵本を村上春樹が翻訳した『あらゆるものがおくりもの』が2026年9月17日に発売(1760円)。人気作家の翻訳仕事が話題の書籍として展開される。","ベストセラー作家のブランド力を『翻訳』という形でも活用する動き。作家名そのものが販売の核となり、著作以外の翻訳・エッセイ・監修などの多角的な出版展開が話題を生む。","村上春樹訳の絵本が9/17発売"+NL+"パトリック・マクドネル作を翻訳"+NL+"作家ブランドの多角活用","人気作家のブランドは著作にとどまらず翻訳・監修・関連本へ拡張され、出版社は作家名を軸にした多角展開で確実な売上を狙う。作家のIP化が進む。","作家ブランドの多角展開(翻訳・監修)","実業之日本社","村上春樹","出版,国内",""],
]
n=13; added=0; skipped=0; bg={}
for row in news:
    u=row[3].strip()
    if u in existing: skipped+=1; continue
    ws.append(row); sr(ws,ws.max_row,n); existing.add(u); added+=1; bg[row[1]]=bg.get(row[1],0)+1
print("BATCH3 added:",added,"skipped:",skipped,"bg:",bg)
wb.save(PATH)
from collections import Counter
wb2=openpyxl.load_workbook(PATH, read_only=True); w=wb2["日次ニュース"]; c=Counter()
for r in w.iter_rows(min_row=2, values_only=True):
    if str(r[0])=="2026-09-14": c[r[1]]+=1
print("TODAY FINAL:",dict(c),"sum",sum(c.values()))
