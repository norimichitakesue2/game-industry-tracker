# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, Alignment
from datetime import date, datetime

TODAY="2026-09-01"
TODAY_D=date(2026,9,1)
PATH="master/ゲーム業界・時事情報収集.xlsx"
YU="Yu Gothic"
def style_row(ws,r,ncol):
    for c in range(1,ncol+1):
        cell=ws.cell(row=r,column=c)
        cell.font=Font(name=YU,size=11)
        cell.alignment=Alignment(wrap_text=True,vertical="top")
def to_d(v):
    if v is None or v=="":return None
    if isinstance(v,datetime):return v.date()
    if isinstance(v,date):return v
    s=str(v).strip().replace("/","-")
    for f in("%Y-%m-%d","%Y-%m"):
        try:return datetime.strptime(s,f).date()
        except:pass
    return None

wb=openpyxl.load_workbook(PATH)

# ---- 横断的問い (8 cols) ----
ws=wb["横断的問い"]
exist=set(str(r[1]).strip() for r in ws.iter_rows(min_row=2,values_only=True) if r[1])
CQ=[
 [TODAY,"配信各社(Apple Music/Spotify)が進めるAI楽曲の“表示義務化・推薦除外”は、ゲームの生成アセットや小説・漫画のAI生成にも同種の開示ルールとして波及するか？",
  "Apple Music「Made With AI」表示義務化／Spotify「AIペルソナ」バッジ","Apple, Spotify, 出版各社",
  "出版(AI小説の開示)、ゲーム(生成アセット表示)、広告・映像の生成AI表記","オープン",
  "音楽が先行する“識別＋隔離”モデルは、AIコンテンツ全般の開示規律の雛形になり得る。ジャンル横断で監視。","横断"],
 [TODAY,"ゲーム機の値上げ疲れ(関税・メモリ高)による消費手控えは、映画館やライブの“体験価格”上昇でも同じ消費者離れを生むか？それとも体験消費は価格耐性が高いか？",
  "米7月ゲーム機支出2020年以来最低／ぴあ総研ライブ市場は単価上昇で最高","ソニー, 任天堂, ぴあ",
  "劇場興行の単価、ライブ単価、サブスク値上げ耐性の比較","オープン",
  "同じ“値上げ”でもハード(所有)は手控え、ライブ(体験)は単価上昇でも動員増。所有と体験で価格弾力性が逆。","横断"],
]
add_cq=0
for r in CQ:
    if str(r[1]).strip() in exist:continue
    ws.append(r);style_row(ws,ws.max_row,8);add_cq+=1
print("横断的問い added:",add_cq)

# ---- メディアミックス追跡 (11 cols) ----
ws=wb["メディアミックス追跡"]
exist=set((str(r[1]).strip(),str(r[4]).strip()) for r in ws.iter_rows(min_row=2,values_only=True) if r[1])
MM=[
 [TODAY,"8月31日のロングサマー","小説/漫画","(原作)","アニメ, 映画","TVアニメ(制作シャフト)＋実写ドラマのW映像化","2026-08-31","未定","発表","シャフト","アニメと実写ドラマを同時展開するW映像化"],
 [TODAY,"高校生家族","漫画","仲間りょう/集英社","映画","実写映画化(香取慎吾主演)","2026-06-25","2027-01-08","制作中","集英社, 東宝","家族全員が高校生。特報公開済"],
 [TODAY,"うるわしの宵の月","漫画","やまもり三香/講談社","映画","実写映画化","(既報)","2026-10-23","公開間近","講談社","道枝駿佑×安斉星来、主題歌なにわ男子"],
]
add_mm=0
for r in MM:
    if (str(r[1]).strip(),str(r[4]).strip()) in exist:continue
    ws.append(r);style_row(ws,ws.max_row,11);add_mm+=1
print("メディアミックス added:",add_mm)

# ---- イベントカレンダー (10 cols) : status recompute + new ----
ws=wb["イベントカレンダー"]
# recompute status for existing rows
upd=0
for row in ws.iter_rows(min_row=2):
    start=to_d(row[0].value); end=to_d(row[1].value); name=row[2].value
    if not name:continue
    st=row[8].value
    if start is None:  # 日程未確定
        continue
    e=end or start
    if e<TODAY_D:new="終了"
    elif start<=TODAY_D<=e:new="開催中"
    else:new="開催予定"
    if st!=new:
        row[8].value=new; row[9].value=TODAY
        for c in row: c.font=Font(name=YU,size=11);c.alignment=Alignment(wrap_text=True,vertical="top")
        upd+=1
print("イベント status updated:",upd)
exist_ev=set(str(r[2]).strip() for r in ws.iter_rows(min_row=2,values_only=True) if r[2])
EV=[
 ["2026-09-03","2026-09-03","State of Play / State of Play Japan(2026年9月)","ショーケース","オンライン","配信","PS向け新作情報。本国版と日本版を同時開催","https://www.eventhubs.com/news/2026/aug/31/state-play-sony-playstation-stream/","開催予定",TODAY],
 ["2026-09-06","2026-09-06","Nintendo Direct(2026年9月)","ショーケース","オンライン","配信","約50分。Switch 2/Switch向け発売予定タイトル中心","https://www.nintendo.com/en-gb/News/Nintendo-Direct/2026/Nintendo-Direct-09-06-2026-3130904.html","開催予定",TODAY],
]
add_ev=0
for r in EV:
    if str(r[2]).strip() in exist_ev:continue
    ws.append(r);style_row(ws,ws.max_row,10);add_ev+=1
print("イベント added:",add_ev)

# ---- 今後起きそうなこと (8 cols) ----
ws=wb["今後起きそうなこと"]
exist=set(str(r[1]).strip() for r in ws.iter_rows(min_row=2,values_only=True) if r[1])
FT=[
 [TODAY,"配信各社のAI楽曲“表示＋推薦除外”が常態化し、AI音楽が実在アーティストの発見面から締め出される→AI楽曲専門の配信網・レーベルが分離し、別市場(BGM/量産用途)として独立していく",
  "中","6ヶ月〜1年","Spotify, Apple, Suno, Udio","Spotify「AIペルソナ」バッジ導入(Spotify newsroom)／Apple Music AI表示義務化(THR)",
  "Spotifyの推薦除外運用後のAI楽曲再生数の変化、AI専門配信・レーベルの登場有無を観測","音楽"],
]
add_ft=0
for r in FT:
    if str(r[1]).strip() in exist:continue
    ws.append(r);style_row(ws,ws.max_row,8);add_ft+=1
print("今後 added:",add_ft)

# ---- ゲームデザイン・トレンド (10 cols) ----
ws=wb["ゲームデザイン・トレンド"]
rows=list(ws.iter_rows(min_row=2))
names={}
for row in rows:
    if row[1].value:names[str(row[1].value).strip()]=row
GD_NAME="レガシー資産のエバーグリーン運用(Mod/リマスターによる旧作長寿命化)"
matched=None
for k,row in names.items():
    if "エバーグリーン" in k or ("レガシー" in k and "長寿命" in k):matched=row;break
if matched:
    matched[3].value="再評価/流行中"
    matched[7].value="新作AAA疲れの中、旧作の作り直し・UGC取り込みで収益寿命を延ばす運用が主流化"
    matched[8].value="NexusMods累計DL250億回・スカイリムSEが半数(Game*Spark 8/29)"
    matched[0].value=TODAY
    for c in matched:c.font=Font(name=YU,size=11);c.alignment=Alignment(wrap_text=True,vertical="top")
    print("ゲームデザイン updated existing")
else:
    ws.append([TODAY,GD_NAME,"運営/構造","再評価/流行中",
      "10年以上前の旧作がMod・リマスター・公式UGC支援で桁違いの寿命を持つ。新作AAAの伸び悩みと対照的に、レガシー作品にユーザー時間が厚く残る",
      "The Elder Scrolls V: Skyrim, The Witcher 3","新作AAAの高コスト・話題化難、Mod文化の成熟、旧作の低価格・高継続性",
      "パブリッシャーはリマスター/公式Mod支援/UGC取り込みで旧IPの収益寿命を延ばす戦略を強める","NexusMods累計DL250億回・スカイリムSEが半数(Game*Spark 8/29)",4])
    style_row(ws,ws.max_row,10)
    print("ゲームデザイン added new")

wb.save(PATH)
print("SAVED aux")
