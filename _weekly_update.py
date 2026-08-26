# -*- coding: utf-8 -*-
import openpyxl

PATH = "master/ゲーム業界・時事情報収集.xlsx"
TODAY = "2026-08-26"
wb = openpyxl.load_workbook(PATH)

deep = wb["掘り下げ・分析"]

deep_rows = [
    [TODAY,
     "改正著作権法「新裁定制度」始動——権利者不明作品×AI時代の権利処理インフラ",
     "2026年4月施行の改正著作権法で新設された「新裁定制度」の実運用がこの夏から本格化。権利者が不明・連絡不能な「未管理著作物」を、文化庁長官の裁定と補償金供託により利用可能にする仕組みで、絶版漫画・過去の楽曲・古い小説の二次利用やデジタルアーカイブ、配信を後押しする。同時に、生成AIの学習・利用を巡る権利処理の枠組みとしても注目される。一方でCODA（スタジオジブリ・任天堂等100社超）や日本動画協会・出版19団体がOpenAI/Sora 2に無断学習停止を求める声明を出すなど、「利活用促進」と「無断利用への抗議」が同時進行。制度が誰の利益に働くかを巡り、権利者団体とAI事業者の綱引きが続く。",
     "文化庁, CODA, スタジオジブリ, 任天堂, 日本動画協会, OpenAI(Sora 2), 出版広報センター",
     "新裁定制度は権利者不明作品の“死蔵”解消と流通拡大が狙いだが、AI学習データの合法確保ルートにもなり得る点が論点。日本の著作権法30条の4はAI学習に寛容とされ、権利者側はオプトアウトや対価還元の欠如に不満を持つ。制度は権利処理の摩擦を制度的に下げる方向に働くため、コンテンツ供給側（出版・音楽・映像）の交渉力とAI事業者の利用範囲のバランスが焦点。",
     "権利処理インフラは漫画・小説・音楽・映像・声のすべてに横断的に効く。AI時代の分配ルールがどのジャンルの慣行を雛形にするかで業界間の力関係が決まる。",
     "https://www.manegy.com/news/detail/15990/ , https://www.bunka.go.jp/",
     "横断"],
    [TODAY,
     "『人間が作った』ことが価値基準に浮上——AI判別不能時代の\"人手証明\"プレミアム",
     "生成AIコンテンツが人間の制作物と判別困難になる中、「人間性」を価値・分配の基準に据える動きが音楽・小説・声の各領域で同時に顕在化。音楽ではDeezerとハンガリーの権利保護団体EJIが2026年3月にライセンス提携し「音楽は人間の創造物であり人間の実演家のみを保護対象とする」としてAI単独生成の録音物にロイヤリティを支払わない方針を採用。小説では生成AI作品が文学賞選考に絡み「AI活用力」と「人手の創作的寄与」の線引きが争点化。声では声優・津田健次郎氏が無断AI音声を巡りTikTok運営を提訴。判別が不能だからこそ「人間が作った証明」が希少価値とルールの基準になりつつある。",
     "Deezer, EJI, 津田健次郎, TikTok(ByteDance), UMG, JASRAC",
     "AIの品質向上で“出力の見分け”が困難になると、価値の源泉は成果物そのものから「誰が・どう作ったか」という来歴（プロビナンス）へ移る。権利団体は分配基準を「人間の関与」に置くことで、AI量産物によるカタログ希薄化から人間の実演家を守ろうとする。これは真贋の技術的検証（電子透かし・開示）と、契約・団体規約による人手保証の二層で進む。",
     "「人手証明」は非AIの差別化を超え、ロイヤリティ分配・受賞資格・出演契約の入場条件になり得る。ジャンル横断の“人間性プレミアム”市場が立ち上がる。",
     "https://www.musicman.co.jp/business/717998 , https://www.morihamada.com/ja/insights/newsletters/131101",
     "横断"],
    [TODAY,
     "gamescom 2026開幕——インディーエリア過去最大(+16%)が示す供給過剰とディスカバリー危機",
     "世界最大級のゲーム見本市gamescom 2026が8月26〜30日、独ケルンで開幕。出展社は67か国1600社超で過去最高、展示スペースは史上初の完全完売、インディーエリアは前年比16%増と最大の伸び。Opening Night LiveでMETRO 2039やWitcher 3新拡張『Songs of the Past』、FF7 Reversal(仮)・無限大ANANTA・天地創造の新展開などが話題を集める。一方でインディーの出展・供給が膨らむほど、Steam上での“発見されにくさ”は深刻化。旧作が収益の大半を占める市場構造（上半期売上の約8割が既存作）と相まって、新規タイトルが埋もれる「供給過剰×ディスカバリー難」が構造課題として浮上している。",
     "gamescom(KoelnMesse), Valve(Steam), CD Projekt RED, Square Enix, NetEase, Deep Silver(METRO)",
     "見本市は“注目の再集約装置”だが、出展増＝新作供給増は個々タイトルの可視性を薄める。Steamの推薦アルゴリズムとウィッシュリスト依存が強まる中、露出は一部の話題作に集中し、ロングテールのインディーは初動を作れない。イベントでの一発露出後に失速する構造が、供給過剰下でより先鋭化する。",
     "インディー供給過剰は漫画（Webtoon量産）・小説（なろう投稿過多）と同型のディスカバリー問題。発見の仕組みを持つ者が価値を握る。",
     "https://www.4gamer.net/games/991/G999110/20260819009/ , https://automaton-media.com/articles/newsjp/20260818-460870/",
     "ゲーム"],
    [TODAY,
     "レーベル公認AI音楽PFへの移行——Udio×UMG新PFとDeezer/EJIの『分配除外』",
     "生成AI音楽が「無断学習・氾濫」から「ライセンス済み生成」へと制度的に整理されるフェーズに入った。UdioはUniversal Music Group(UMG)との2025年10月の和解を受けダウンロード機能を停止、PF内での生成・共有モデルへ移行し、2026年にUMGと共同のライセンス済み新プラットフォームを立ち上げ予定。配信側ではDeezerがEJIと2026年3月にライセンス提携し、AI単独生成の録音物を分配対象から除外する方針を採用。「AI楽曲を排除」ではなく「権利処理済みAI」と「人間の実演」を仕分けて共存させる設計が主流化しつつある。カタログ希薄化とロイヤリティ原資の奪い合いへの、業界の制度的回答が形になってきた。",
     "Udio, Universal Music Group, Deezer, EJI, Suno",
     "ストリーミングでは消費者支払いに占めるテック/PFの取り分が1999年の4%から2023年に33%へ拡大し、アーティスト取り分は14%で横ばい。原資が限られる中でAI量産曲が再生数を吸えば人間アーティストの分配はさらに薄まる。ゆえに「ライセンス済みAI」への囲い込みと「非ライセンスAIの分配除外」が原資防衛の二本柱になる。",
     "レーベル公認AIモデルは、Sora 2を巡る映像・声優の権利交渉の先行事例。音楽の分配ルールが他ジャンルの雛形になり得る。",
     "https://www.musicman.co.jp/business/717998 , https://ai-revolution.co.jp/media/ai-in-music/",
     "音楽"],
    [TODAY,
     "Webtoon『創作性不足』批判とファン層の乖離——量産モデルの質的踊り場",
     "縦読みWebtoon市場は世界的には高成長（プラットフォーム市場でCAGR30%級の予測）が続く一方、質と読者満足を巡る不協和音が表面化。読者調査では「Webtoonは創作性が足りない」との声が目立ち、量産・スタジオ分業・アルゴリズム最適化で似通った作品が増えたことへの不満と、コアなマンガ読者層との乖離が指摘される。産業側でも最大手NAVER系Webtoon Entertainmentが従業員3%削減、傘下Wattpadも15%削減、カカオエンタも希望退職・仏撤退と、成長の裏で再編が進む。数量成長が頭打ちに近づく中、「量」から「質・独自性」への転換が次の競争軸として問われ始めている。",
     "NAVER Webtoon(Webtoon Entertainment), Wattpad, カカオエンターテインメント, LINEマンガ, ピッコマ",
     "Webtoonのスタジオ量産モデルは供給量とローンチ本数で市場を拡大したが、テンプレ化・素材の使い回しが“創作性”の希薄化を招く。アルゴリズム配信は初速のある作品に露出を集中させ、実験的作品が育ちにくい。結果、ライトユーザーは伸ばせても既存マンガ読者の熱量を取り込めず、ファン層が分断される。",
     "「量産×アルゴリズム配信＝独自性の希薄化」はゲーム（インディー供給過剰）や小説（なろう量産）と同型。質の担保が次の分岐点。",
     "https://premium.kai-you.net/article/892 , https://gendai.media/articles/-/143135",
     "漫画"],
]

existing_deep = set()
for r in deep.iter_rows(min_row=2, values_only=True):
    if r[1]:
        existing_deep.add(str(r[1]).strip())

added = []
for row in deep_rows:
    if row[1].strip() in existing_deep:
        print("SKIP dup:", row[1]); continue
    deep.append(row); added.append(row[1])
print("APPENDED deep-dive:", len(added))

def update_trend(sheet_name, theme, genkyo, doko, kigyo, yosoku, kizuki, link):
    ws = wb[sheet_name]
    for row in ws.iter_rows(min_row=2):
        if row[1].value and str(row[1].value).strip() == theme:
            row[0].value=TODAY; row[3].value=genkyo; row[4].value=doko
            row[5].value=kigyo; row[6].value=yosoku; row[7].value=kizuki; row[8].value=link
            print("UPDATED", sheet_name, "::", theme); return True
    print("NOTFOUND", sheet_name, "::", theme); return False

update_trend("業界トレンド・技術","ゲームショウ／トレードショウの再集約",
    "gamescom 2026(8/26-30・ケルン)が出展1600社超で過去最高、展示は史上初の完全完売。インディーエリアは前年比16%増と最大の伸び。",
    "Opening Night LiveでMETRO 2039、Witcher 3新拡張『Songs of the Past』、FF7 Reversal(仮)・無限大ANANTA・天地創造の新展開などを一挙披露。物理見本市の集客力が回復・過熱。",
    "KoelnMesse, Valve, CD Projekt RED, Square Enix, NetEase, Deep Silver",
    "出展増＝新作供給増でSteam上のディスカバリー難が深刻化。イベント露出の一極集中とロングテール埋没が並行進行。",
    "見本市の“再集約”とインディー供給過剰は表裏。発見の仕組み（推薦・ウィッシュリスト）を握る者が価値を左右。",
    "https://www.4gamer.net/games/991/G999110/20260819009/")

update_trend("【音楽】トレンド","生成AIと音楽制作・声の権利",
    "生成AI音楽が「無断学習・氾濫」から「ライセンス済み生成」へ移行。UdioはUMGとの和解後DL停止＋PF内生成モデルに転換、2026年にUMG共同の公認PFを予定。",
    "DeezerがEJIと提携(2026/3)しAI単独生成の録音物を分配対象から除外。「人間の実演家のみ保護」を分配基準に。声優・津田健次郎氏はAI音声巡りTikTokを提訴。",
    "Udio, UMG, Deezer, EJI, Suno, TikTok",
    "「公認AI」と「人間の実演」を仕分けて共存させる設計が主流化。非ライセンスAI曲は分配・掲載から段階的に排除へ。",
    "音楽の分配ルール設計が、Sora 2を巡る映像・声優交渉の先行事例になる。",
    "https://www.musicman.co.jp/business/717998")

update_trend("【音楽】トレンド","ストリーミング収益分配の議論",
    "消費者支払いに占めるテック/PF取り分は1999年4%→2023年33%へ拡大、アーティスト取り分は14%で横ばい（Musicman 2026報告）。",
    "原資が限られる中でAI量産曲が再生数を吸えば人間アーティストの分配がさらに希薄化。権利団体は「人間の関与」を分配基準化する動き。",
    "Spotify, Apple Music, UMG, Deezer, CPRA",
    "実演家への適切報酬を求める海外の圧力（CPRA等）と連動し、分配モデル見直し・AI曲の扱い規定が進む。",
    "「成長しても増えないアーティスト取り分」は楽曲カタログ証券化・ファンド化の裏返し。原資の奪い合いが激化。",
    "https://www.cpra.jp/cpra_article/article/000714.html")

update_trend("【漫画】トレンド","縦読みWebtoonの浸透",
    "世界市場は高成長予測（PF市場でCAGR30%級）が続く一方、「創作性が足りない」との読者批判とコア読者層の乖離が表面化。",
    "NAVER系Webtoon Entertainmentが3%削減、Wattpad15%削減、カカオエンタも希望退職・仏撤退。量産・スタジオ分業のテンプレ化に不満が集中。",
    "NAVER Webtoon, Wattpad, カカオエンタ, LINEマンガ, ピッコマ",
    "数量成長の頭打ちが近づき「量」から「質・独自性」への競争軸転換。実験作を育てる編集・発見の仕組みが差別化要因に。",
    "量産×アルゴリズム配信＝独自性希薄化はゲーム・小説と同型のディスカバリー問題。",
    "https://premium.kai-you.net/article/892")

wb.save(PATH)
print("SAVED", PATH)
